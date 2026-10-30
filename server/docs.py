# -*- coding: utf-8 -*-
"""文档读取：把 Word / PDF / Excel / PPT / 网页 / 电子书拆成字。

为什么要单独一层：`fs_read` 只会读字节，而用户拖进来的是 docx、是 PDF、
是老师发的课件。二进制被原样当文本读出来是一堆乱码 —— 那不是「读不了」，
是「看起来读了」。所以这里按格式拆，拆不出来就明说拆不出来。

拆完统一返回纯文本 + 一句「这是什么、有多少字」，调用方负责分段给模型看。
"""
import html as _html
import io
import os
import re
import zipfile

# 纯文本：不用解析，按行读就行（fsops 自己处理）
TEXT_EXT = {
    ".txt", ".md", ".markdown", ".json", ".jsonl", ".yaml", ".yml", ".csv", ".tsv",
    ".log", ".ini", ".cfg", ".conf", ".toml", ".xml", ".html", ".htm", ".py", ".js",
    ".ts", ".tsx", ".jsx", ".java", ".c", ".cpp", ".h", ".hpp", ".cs", ".go", ".rs",
    ".php", ".rb", ".sql", ".sh", ".bat", ".ps1", ".css", ".scss", ".vue", ".kt",
    ".swift", ".lua", ".pl", ".r", ".m", ".sql", ".env", ".gitignore", ".properties",
    ".srt", ".ass", ".vtt", ".tex", ".rst", ".org", ".srt",
}

# 要拆的文档格式
DOC_EXT = {
    ".docx": "docx", ".docm": "docx",
    ".doc": "doc",                      # 老二进制 .doc，尽力而为
    ".pdf": "pdf",
    ".pptx": "pptx", ".ppsx": "pptx",
    ".xlsx": "xlsx", ".xlsm": "xlsx",
    ".xls": "xls", ".ppt": "ppt",          # 老二进制格式，只报「另存一下」
    ".wps": "doc", ".et": "xls", ".dps": "ppt",   # WPS 三件套其实同源
    ".pages": "pages", ".numbers": "pages", ".key": "pages",
    ".odt": "odt", ".ods": "ods", ".odp": "odp",
    ".rtf": "rtf",
    ".epub": "epub",
    ".mobi": "mobi",
}

# 只报身份，不拆内容：视频/音频/图片/压缩包/程序
MEDIA_EXT = {
    ".mp4": "视频", ".mkv": "视频", ".webm": "视频", ".avi": "视频", ".mov": "视频",
    ".flv": "视频", ".wmv": "视频", ".m4v": "视频", ".ts": "视频",
    ".mp3": "音频", ".wav": "音频", ".flac": "音频", ".m4a": "音频", ".aac": "音频",
    ".ogg": "音频", ".wma": "音频",
    ".png": "图片", ".jpg": "图片", ".jpeg": "图片", ".gif": "图片", ".webp": "图片",
    ".bmp": "图片", ".ico": "图片", ".tif": "图片", ".tiff": "图片", ".psd": "图片",
    ".zip": "压缩包", ".rar": "压缩包", ".7z": "压缩包", ".tar": "压缩包", ".gz": "压缩包",
    ".exe": "程序", ".dll": "程序", ".msi": "程序", ".apk": "程序", ".bin": "程序",
    ".db": "数据库", ".sqlite": "数据库", ".sqlite3": "数据库",
    ".onnx": "模型", ".pt": "模型", ".safetensors": "模型", ".gguf": "模型",
}


def kind_of(path):
    ext = os.path.splitext(path)[1].lower()
    if ext in TEXT_EXT:
        return "text"
    if ext in DOC_EXT:
        return DOC_EXT[ext]
    if ext in MEDIA_EXT:
        return "media"
    return "unknown"


# ------------------------------------------------------------------ 各格式
def _docx(path):
    import docx
    d = docx.Document(path)
    out = []
    for p in d.paragraphs:
        t = (p.text or "").strip()
        if t:
            out.append(t)
    for ti, tb in enumerate(d.tables, 1):
        rows = []
        for row in tb.rows:
            cells = [(c.text or "").replace("\n", " ").strip() for c in row.cells]
            if any(cells):
                rows.append(" | ".join(cells))
        if rows:
            out.append("【表 %d】" % ti)
            out.extend(rows)
    return "\n".join(out)


def _pptx(path):
    from pptx import Presentation
    prs = Presentation(path)
    out = []
    for i, slide in enumerate(prs.slides, 1):
        buf = []
        for sh in slide.shapes:
            if getattr(sh, "has_text_frame", False):
                t = (sh.text_frame.text or "").strip()
                if t:
                    buf.append(t)
            if getattr(sh, "has_table", False):
                for row in sh.table.rows:
                    cells = [(c.text or "").replace("\n", " ").strip() for c in row.cells]
                    if any(cells):
                        buf.append(" | ".join(cells))
        try:
            nf = slide.notes_slide.notes_text_frame
            if nf and (nf.text or "").strip():
                buf.append("[备注] " + nf.text.strip())
        except Exception:
            pass
        if buf:
            out.append("—— 第 %d 页 ——" % i)
            out.extend(buf)
    return "\n".join(out)


def _xlsx(path, max_rows=600):
    import openpyxl
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    out = []
    for ws in wb.worksheets:
        out.append("—— 工作表 %s ——" % (ws.title or "?"))
        n = 0
        for row in ws.iter_rows(values_only=True):
            cells = ["" if v is None else str(v).replace("\n", " ").strip() for v in row]
            while cells and cells[-1] == "":
                cells.pop()
            if not any(cells):
                continue
            out.append(" | ".join(cells))
            n += 1
            if n >= max_rows:
                out.append("（这张表只取了前 %d 行）" % max_rows)
                break
    wb.close()
    return "\n".join(out)


def _pdf(path, max_pages=200):
    """先 pdfplumber（排版好），不行再 pdfminer（快但糙）。"""
    text = ""
    total = 0
    try:
        import pdfplumber
        pages = []
        with pdfplumber.open(path) as pdf:
            total = len(pdf.pages)
            for i, pg in enumerate(pdf.pages[:max_pages], 1):
                t = pg.extract_text() or ""
                if t.strip():
                    pages.append("—— 第 %d 页 ——\n%s" % (i, t.strip()))
        text = "\n".join(pages)
        if total > max_pages:
            text += "\n（这份 PDF 共 %d 页，只拆了前 %d 页）" % (total, max_pages)
    except Exception:
        text = ""
    if not text.strip():
        try:
            from pdfminer.high_level import extract_text
            text = extract_text(path, maxpages=max_pages) or ""
        except Exception:
            text = ""
    return text


def _odf(path):
    """odt/ods/odp 都是 zip 里的 content.xml，取段落文字就行。"""
    with zipfile.ZipFile(path) as z:
        try:
            raw = z.read("content.xml")
        except KeyError:
            return ""
    from lxml import etree
    root = etree.fromstring(raw)
    out = []
    for el in root.iter():
        tag = etree.QName(el).localname if isinstance(el.tag, str) else ""
        if tag in ("p", "h", "table-cell"):
            t = "".join(el.itertext()).strip()
            if t:
                out.append(t)
    return "\n".join(out)


def _epub(path):
    from bs4 import BeautifulSoup
    out = []
    with zipfile.ZipFile(path) as z:
        names = [n for n in z.namelist()
                 if n.lower().endswith((".xhtml", ".html", ".htm"))]
        for n in sorted(names):
            try:
                soup = BeautifulSoup(z.read(n), "lxml")
            except Exception:
                continue
            for bad in soup(["script", "style"]):
                bad.decompose()
            t = soup.get_text("\n")
            t = re.sub(r"\n{2,}", "\n", t).strip()
            if t:
                out.append(t)
    return "\n\n".join(out)


_RTF_CTRL = re.compile(r"\\([a-zA-Z]+)(-?\d+)? ?")


def _rtf_ctrl(m):
    w = m.group(1)
    if w in ("par", "line", "row"):
        return "\n"
    if w == "tab":
        return "\t"
    return ""


def _strip_group(s, name):
    """删掉 {\fonttbl ...} 这类整组。要数括号 —— 它是嵌套的，
    正则的 [^{}]* 只够到最里层。"""
    key = "{\\" + name
    while True:
        i = s.find(key)
        if i < 0:
            return s
        depth, j = 0, i
        while j < len(s):
            if s[j] == "{":
                depth += 1
            elif s[j] == "}":
                depth -= 1
                if depth == 0:
                    j += 1
                    break
            j += 1
        s = s[:i] + " " + s[j:]


def _rtf(path, max_bytes=8 * 1024 * 1024):
    """RTF 就手拆：去掉字体表/颜色表和所有控制字，剩下的就是字。

    两个坑都踩过：
      1. 先 `replace("\\par", "\\n")` 会把 `\\pard` 也切开，剩一个孤零零的 d。
         所以控制字必须一次性按「完整单词」匹配，不能拿子串去替。
      2. 中文有两种写法：`\\'ba\\'af`（GBK 字节）和 `\\u20989?`（Unicode）。
         前者要拼成字节再按 gbk 解，后者本来就是字符 —— 如果混在一起做
         `encode("latin-1")`，汉字会被 ignore 掉。所以 `\\u` 先换成占位符，
         等字节层处理完再换回来。
    """
    raw = open(path, "rb").read(max_bytes)
    s = raw.decode("latin-1", "ignore")

    uni = []

    def _hold(m):
        uni.append(chr(int(m.group(1)) % 65536))
        return "\x01%d\x02" % (len(uni) - 1)

    s = re.sub(r"\\u(-?\d+)\??", _hold, s)                       # \uN? 先存起来
    s = re.sub(r"\\'([0-9a-fA-F]{2})", lambda m: chr(int(m.group(1), 16)), s)
    for g in ("fonttbl", "colortbl", "stylesheet", "info", "pict", "object",
              "header", "footer", "listtable", "listoverridetable"):
        s = _strip_group(s, g)
    s = re.sub(r"\{\\\*.*?\}", " ", s, flags=re.S)               # 目标块
    s = _RTF_CTRL.sub(_rtf_ctrl, s)
    s = s.replace("{", "").replace("}", "")
    # 字节层：latin-1 铺开的字符按 gbk 收回去（中文 RTF 基本是这个码页）
    try:
        s = s.encode("latin-1", "ignore").decode("gbk", "ignore")
    except Exception:
        pass
    s = re.sub(r"\x01(\d+)\x02", lambda m: uni[int(m.group(1))], s)   # 换回 Unicode 字
    s = re.sub(r"[ \t]{2,}", " ", s)
    return re.sub(r"\n{3,}", "\n\n", s).strip()


def _doc(path):
    """老 .doc（OLE 二进制）。没有 olefile 就只能说清楚做不到，
    不要拿乱码冒充「读到了」。"""
    try:
        import olefile
    except ImportError:
        return ("__need__", "这是旧版 .doc（二进制格式）。本机没装 olefile，拆不开。"
                           "用 Word 另存成 .docx 再给我，或者直接把内容粘过来。")
    try:
        if not olefile.isOleFile(path):
            return ("__need__", "这个 .doc 不是 OLE 格式，认不出来。另存成 .docx 试试。")
        ole = olefile.OleFileIO(path)
        if not ole.exists("WordDocument"):
            return ("__need__", "OLE 里没有 WordDocument 流，这不像 Word 文件。")
        data = ole.openstream("WordDocument").read()
        ole.close()
        # 粗暴取可打印片段：中文按 utf-16le 找
        s = data.decode("utf-16-le", "ignore")
        s = "".join(ch if (ch.isprintable() or ch in "\n\t") else " " for ch in s)
        s = re.sub(r" {2,}", " ", s)
        return re.sub(r"\n{3,}", "\n\n", s).strip()
    except Exception as e:
        return ("__need__", "拆 .doc 的时候出错：%s。另存成 .docx 最稳。" % e)


def _html_text(path):
    from bs4 import BeautifulSoup
    raw = open(path, "rb").read(4 * 1024 * 1024)
    soup = BeautifulSoup(raw, "lxml")
    for bad in soup(["script", "style", "noscript"]):
        bad.decompose()
    t = soup.get_text("\n")
    return re.sub(r"\n{3,}", "\n\n", t).strip()


# ------------------------------------------------------------------ 入口
def extract(path, max_chars=200000):
    """拆一份文档。返回 {ok, kind, text, chars, note}；拆不动就说拆不动。"""
    kind = kind_of(path)
    ext = os.path.splitext(path)[1].lower()
    size = os.path.getsize(path)
    if kind == "text" or kind == "unknown":
        return {"ok": False, "kind": kind, "error": "这走的是纯文本那条路，不该进这里"}
    if kind == "media":
        return {"ok": False, "kind": "media", "media": MEDIA_EXT.get(ext, "二进制"),
                "size": size,
                "error": "%s 文件（%.1fMB），不拆内容。路径已经给你了，要用它就说路径"
                         % (MEDIA_EXT.get(ext, "二进制"), size / 1048576.0)}
    try:
        if kind == "docx":
            text = _docx(path)
        elif kind == "pptx":
            text = _pptx(path)
        elif kind == "xlsx":
            text = _xlsx(path)
        elif kind == "pdf":
            text = _pdf(path)
        elif kind in ("xls", "ppt"):
            return {"ok": False, "kind": kind,
                    "error": "旧版 .%s 是二进制格式，本机拆不开（要 Office 的私有解码器）。"
                             "用 Office/WPS 另存成 .%sx 就能读。" % (kind, kind)}
        elif kind == "pages":
            return {"ok": False, "kind": "pages",
                    "error": "苹果 iWork 的包格式，拆不开。导出成 PDF 或 docx 就能读。"}
        elif kind == "odt" or kind in ("ods", "odp"):
            text = _odf(path)
        elif kind == "epub":
            text = _epub(path)
        elif kind == "rtf":
            text = _rtf(path)
        elif kind == "doc":
            r = _doc(path)
            if isinstance(r, tuple):
                return {"ok": False, "kind": "doc", "error": r[1], "need": "docx"}
            text = r
        elif kind == "mobi":
            return {"ok": False, "kind": "mobi",
                    "error": "mobi 是亚马逊的私有格式，本机没有解析器。转成 epub 就能读。"}
        else:
            return {"ok": False, "kind": kind, "error": "这个格式还不认识：%s" % ext}
    except ImportError as e:
        return {"ok": False, "kind": kind,
                "error": "缺解析库（%s）。这台机器没装，装一下就能读。" % e}
    except Exception as e:
        return {"ok": False, "kind": kind,
                "error": "%s 这份文件拆不开：%s。可能加密了，或者文件本身坏了。"
                         % (ext, str(e)[:160])}

    text = (text or "").replace("\r\n", "\n").replace("\r", "\n")
    text = re.sub(r"\n{4,}", "\n\n\n", text).strip()
    if not text:
        return {"ok": False, "kind": kind, "size": size,
                "error": "%s 拆出来是空的 —— 可能是扫描件（整页是图，没有文字层），"
                         "也可能是空文档。" % ext}
    chars = len(text)
    cut = chars > max_chars
    return {"ok": True, "kind": kind, "chars": chars, "truncated": cut,
            "text": text[:max_chars],
            "note": ("原文 %d 字，这里只取了前 %d 字" % (chars, max_chars)) if cut else ""}
